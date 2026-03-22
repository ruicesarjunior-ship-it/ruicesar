"""
Analisador de procedimentos extrajudiciais do MPBA.

Fluxo:
    1. Extrai texto do arquivo (DOCX, PDF ou TXT)
    2. Chama Claude API → análise jurídica + decisão de arquivamento
    3. Se viável, extrai zonas vermelhas do template escolhido
    4. Chama Claude API novamente → conteúdo para cada zona vermelha
    5. Aplica tudo ao template e gera o .docx final
"""
from __future__ import annotations

import json
import logging
import re
import time
import shutil
from datetime import datetime
from pathlib import Path
from typing import Any

import anthropic

from config import settings
from scripts.office.editor import apply_analysis_to_template, extract_red_zones
from scripts.office.validate import validate

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Extração de texto do documento de entrada
# ---------------------------------------------------------------------------

def extract_text(filepath: str | Path) -> str:
    """Extrai texto de .docx, .pdf ou .txt."""
    filepath = Path(filepath)
    suffix = filepath.suffix.lower()

    if suffix == ".docx":
        return _extract_docx(filepath)
    elif suffix == ".pdf":
        return _extract_pdf(filepath)
    elif suffix in (".txt", ".text", ".md"):
        return filepath.read_text(encoding="utf-8", errors="ignore")
    else:
        raise ValueError(f"Formato não suportado: {suffix}. Use .docx, .pdf ou .txt")


def _extract_docx(filepath: Path) -> str:
    from docx import Document  # type: ignore
    doc = Document(str(filepath))
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)


def _extract_pdf(filepath: Path) -> str:
    import fitz  # PyMuPDF  # type: ignore
    doc = fitz.open(str(filepath))
    parts = []
    for page in doc:
        text = page.get_text()
        if text.strip():
            parts.append(text)
    doc.close()
    return "\n".join(parts)


# ---------------------------------------------------------------------------
# Carregamento do system prompt
# ---------------------------------------------------------------------------

def _load_system_prompt() -> str:
    prompt_path = Path(__file__).parent / "prompts" / "system_prompt.md"
    return prompt_path.read_text(encoding="utf-8")


# ---------------------------------------------------------------------------
# Chamada à API do Claude — Fase 1: Análise
# ---------------------------------------------------------------------------

def analyze_document(document_text: str) -> dict[str, Any]:
    """
    Chama Claude com o texto do documento e retorna a análise estruturada.
    Retorna dict com as chaves definidas no system_prompt.md.
    """
    client = anthropic.Anthropic(api_key=settings.anthropic_api_key)
    system_prompt = _load_system_prompt()

    # Limita o texto para respeitar o limite de tokens da API
    texto_limitado = document_text[:12000]
    if len(document_text) > 12000:
        logger.info("Documento truncado para 12.000 caracteres (original: %d)", len(document_text))

    user_message = (
        "Analise o procedimento extrajudicial abaixo e retorne o JSON conforme instruído "
        "no system prompt.\n\n"
        "=== DOCUMENTO ===\n\n"
        f"{texto_limitado}\n\n"
        "=== FIM DO DOCUMENTO ==="
    )

    logger.info("Enviando documento para análise (Claude API)...")

    response = client.messages.create(
        model=settings.claude_model,
        max_tokens=4096,
        system=system_prompt,
        messages=[{"role": "user", "content": user_message}],
    )

    raw_text = response.content[0].text
    return _parse_json_response(raw_text)


# ---------------------------------------------------------------------------
# Chamada à API do Claude — Fase 2: Preenchimento das zonas vermelhas
# ---------------------------------------------------------------------------

def generate_zone_content(
    document_text: str,
    analysis: dict[str, Any],
    zones: list[dict],
) -> dict[int, str]:
    """
    Dado o documento e a análise, pede ao Claude que preencha cada zona vermelha
    do template com o conteúdo apropriado.
    Retorna dict {zone_id: texto_substituto}.
    """
    if not zones:
        return {}

    client = anthropic.Anthropic(api_key=settings.anthropic_api_key)

    # Monta a listagem de zonas para o Claude
    zones_description = "\n".join(
        f"  ZONA {z['zone_id']}: "
        f"[antes: \"{z['context_before'][:80]}\"] "
        f"[VERMELHO: \"{z['red_text'][:120]}\"] "
        f"[depois: \"{z['context_after'][:80]}\"]"
        for z in zones
    )

    hipotese = analysis.get("analise", {}).get("hipotese_recomendada", "")
    modelo = analysis.get("analise", {}).get("modelo_word", "")
    conteudo = analysis.get("conteudo", {})

    system = (
        "Você é um Promotor de Justiça do MPBA especialista em promoções de arquivamento. "
        "Você preencherá os campos variáveis (ZONAS VERMELHAS) de um template de promoção de "
        "arquivamento. JAMAIS invente informações não presentes no documento. Use "
        "[INSERIR: descrição] para dados faltantes. "
        "Retorne APENAS um JSON válido: {\"replacements\": {\"0\": \"texto\", \"1\": \"texto\", ...}}"
    )

    user_message = (
        f"Hipótese selecionada: {hipotese}\n"
        f"Modelo Word: {modelo}\n\n"
        f"Conteúdo já gerado pela análise:\n{json.dumps(conteudo, ensure_ascii=False, indent=2)}\n\n"
        f"=== ZONAS VERMELHAS DO TEMPLATE ===\n{zones_description}\n\n"
        f"=== DOCUMENTO ORIGINAL ===\n{document_text[:3000]}\n\n"
        "Para cada ZONA, forneça o texto substituto adequado ao caso concreto. "
        "Use o conteúdo já gerado pela análise onde aplicável. "
        "Para zonas de narrativa/relatório, use o texto do documento. "
        "Retorne JSON: {\"replacements\": {\"<zone_id>\": \"<texto>\", ...}}"
    )

    logger.info("Gerando conteúdo para %d zonas vermelhas...", len(zones))
    logger.info("Aguardando 15s para respeitar limite de velocidade da API...")
    time.sleep(15)

    response = client.messages.create(
        model=settings.claude_model,
        max_tokens=4096,
        system=system,
        messages=[{"role": "user", "content": user_message}],
    )

    raw = response.content[0].text
    parsed = _parse_json_response(raw)
    raw_replacements = parsed.get("replacements", {})

    # Converte chaves para int
    return {int(k): v for k, v in raw_replacements.items()}


# ---------------------------------------------------------------------------
# Parser de JSON da resposta do Claude
# ---------------------------------------------------------------------------

def _parse_json_response(text: str) -> dict[str, Any]:
    """Extrai e parseia o JSON do texto de resposta do Claude."""
    # Tenta extrair bloco ```json ... ```
    match = re.search(r"```(?:json)?\s*(\{.*?\})\s*```", text, re.DOTALL)
    if match:
        json_str = match.group(1)
    else:
        # Tenta encontrar { ... } diretamente
        match = re.search(r"(\{.*\})", text, re.DOTALL)
        if match:
            json_str = match.group(1)
        else:
            raise ValueError(f"Nenhum JSON encontrado na resposta do Claude:\n{text[:500]}")

    try:
        return json.loads(json_str)
    except json.JSONDecodeError as e:
        raise ValueError(f"JSON inválido na resposta do Claude: {e}\n\nJSON recebido:\n{json_str[:500]}") from e


# ---------------------------------------------------------------------------
# Montagem do analysis dict para o editor
# ---------------------------------------------------------------------------

def _build_editor_analysis(
    analysis_response: dict[str, Any],
    zone_content: dict[int, str],
    zones: list[dict],
    cidade: str,
) -> dict[str, Any]:
    """Converte a resposta do Claude para o formato esperado pelo editor."""
    analise = analysis_response.get("analise", {})
    named_vars = analysis_response.get("named_vars", {})
    conteudo = analysis_response.get("conteudo", {})
    blocos = analysis_response.get("blocos_condicionais", {})

    # Completa named_vars com dados fixos
    hoje = datetime.now()
    meses = [
        "janeiro", "fevereiro", "março", "abril", "maio", "junho",
        "julho", "agosto", "setembro", "outubro", "novembro", "dezembro",
    ]
    data_extenso = (
        named_vars.get("dataExtenso")
        or f"{cidade}, {hoje.day} de {meses[hoje.month - 1]} de {hoje.year}"
    )

    named_vars.setdefault("cidadeOrgaoUnidade", cidade)
    named_vars["dataExtenso"] = data_extenso
    named_vars.setdefault("membro", "RUI CÉSAR FARIAS DOS SANTOS JÚNIOR")

    extra_replacements: dict[str, str] = {
        "MARCO AURÉLIO RUBICK DA SILVA": "RUI CÉSAR FARIAS DOS SANTOS JÚNIOR",
        **(analysis_response.get("extra_named_replacements") or {}),
    }

    keep_markers = analysis_response.get("keep_hypothesis_markers", [])

    # Mescla zone_content com conteúdo já gerado na análise
    # Se não há zonas mapeadas explicitamente, usa o conteúdo da análise
    final_zone_replacements = dict(zone_content)

    # Para zonas não preenchidas pelo segundo call, tenta mapear pelo contexto
    if zones and conteudo:
        for zone in zones:
            zid = zone["zone_id"]
            if zid in final_zone_replacements:
                continue
            ctx = (zone["context_before"] + " " + zone["context_after"]).lower()
            # Tenta mapear contexto → campo de conteúdo
            mapping_rules = [
                (["objeto", "portaria", "instaurado", "apurar"], "paragrafo_objeto"),
                (["narrativa", "fatos", "relatório", "diligências"], "paragrafo_narrativa"),
                (["transição", "perspectiva", "viabilizar"], "paragrafo_transicao"),
                (["aplicação", "caso concreto", "anos", "lapso"], "paragrafo_aplicacao"),
                (["prescrição", "prescricional", "prazo"], "paragrafo_prescricao"),
                (["ressarcimento", "erário", "dano"], "paragrafo_ressarcimento"),
            ]
            for keywords, field in mapping_rules:
                if any(kw in ctx for kw in keywords):
                    value = conteudo.get(field)
                    if value:
                        final_zone_replacements[zid] = value
                        break

    return {
        "zone_replacements": final_zone_replacements,
        "named_vars": named_vars,
        "keep_hypothesis_markers": keep_markers,
        "extra_named_replacements": extra_replacements,
        "_zones": zones,
        "conteudo": conteudo,
    }


# ---------------------------------------------------------------------------
# Processamento principal de um arquivo
# ---------------------------------------------------------------------------

def process_file(filepath: str | Path) -> dict[str, Any]:
    """
    Processa um arquivo de procedimento extrajudicial.

    Retorna dict:
        status   : "sucesso" | "nao_viavel" | "modelo_ausente" | "erro"
        output   : Path do .docx gerado (se sucesso)
        analysis : dict com a análise
        message  : str descritivo
    """
    filepath = Path(filepath)
    logger.info("Processando: %s", filepath.name)

    try:
        # 1. Extrai texto
        document_text = extract_text(filepath)
        if not document_text.strip():
            return {"status": "erro", "message": "Arquivo vazio ou sem texto extraível"}

        # 2. Análise jurídica
        analysis_response = analyze_document(document_text)
        analise = analysis_response.get("analise", {})

        if not analise.get("viavel_arquivamento"):
            reason = analise.get("razao_nao_viavel", "Nenhuma hipótese de arquivamento identificada.")
            logger.info("Arquivamento não viável: %s", reason)
            _save_analysis_report(filepath, analysis_response, viavel=False)
            return {
                "status": "nao_viavel",
                "analysis": analysis_response,
                "message": reason,
            }

        modelo_key = analise.get("modelo_word")
        if not modelo_key:
            return {"status": "erro", "message": "Claude não identificou o modelo Word a usar."}

        template_path = settings.get_template_path(modelo_key)
        if not template_path:
            msg = (
                f"Template '{modelo_key}' não encontrado em '{settings.pasta_modelos}'. "
                f"Adicione o arquivo .docx correspondente à pasta 'modelos/'."
            )
            logger.warning(msg)
            _save_analysis_report(filepath, analysis_response, viavel=True)
            return {"status": "modelo_ausente", "analysis": analysis_response, "message": msg}

        # 3. Extrai zonas vermelhas do template
        doc_xml = _get_doc_xml_path(template_path)
        zones = extract_red_zones(doc_xml)
        logger.info("Template '%s': %d zonas vermelhas encontradas.", template_path.name, len(zones))

        # 4. Gera conteúdo para as zonas
        zone_content = generate_zone_content(document_text, analysis_response, zones)

        # 5. Monta analysis dict para o editor
        editor_analysis = _build_editor_analysis(
            analysis_response, zone_content, zones, cidade=settings.cidade_orgao
        )

        # 6. Gera o documento
        output_path = _build_output_path(filepath, analise)
        apply_analysis_to_template(template_path, output_path, editor_analysis)

        # 7. Valida
        val = validate(output_path)
        if not val["ok"]:
            logger.warning("Validação com erros: %s", val["errors"])
        else:
            logger.info("Documento gerado e validado: %s", output_path)

        _save_analysis_report(filepath, analysis_response, viavel=True, output_path=output_path)

        return {
            "status": "sucesso",
            "output": output_path,
            "analysis": analysis_response,
            "message": f"Documento gerado: {output_path.name}",
            "validation": val,
        }

    except Exception as exc:
        logger.exception("Erro ao processar '%s': %s", filepath.name, exc)
        return {"status": "erro", "message": str(exc)}


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _get_doc_xml_path(template_path: Path) -> Path:
    """Extrai temporariamente o document.xml para leitura das zonas."""
    import zipfile, tempfile
    tmp = Path(tempfile.mkdtemp(prefix="mpba_read_"))
    with zipfile.ZipFile(template_path, "r") as zf:
        zf.extract("word/document.xml", tmp)
    return tmp / "word" / "document.xml"


def _build_output_path(input_file: Path, analise: dict) -> Path:
    settings.pasta_saida.mkdir(parents=True, exist_ok=True)
    tipo = analise.get("tipo_procedimento", "PROC")
    numero = analise.get("numero_idea", "SEM_NUMERO").replace("/", "-").replace(" ", "_")
    hipotese = analise.get("hipotese_recomendada", "")
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    stem = f"ARQUIVAMENTO_{tipo}_{numero}_{hipotese}_{timestamp}"
    return settings.pasta_saida / f"{stem}.docx"


def _save_analysis_report(
    input_file: Path,
    analysis: dict,
    viavel: bool,
    output_path: Path | None = None,
) -> None:
    """Salva um relatório .txt da análise ao lado do output."""
    settings.pasta_saida.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    report_path = settings.pasta_saida / f"ANALISE_{input_file.stem}_{timestamp}.txt"

    analise = analysis.get("analise", {})
    lines = [
        "=" * 70,
        "RELATÓRIO DE ANÁLISE — PIPELINE MPBA",
        "=" * 70,
        f"Arquivo de entrada  : {input_file.name}",
        f"Data/hora           : {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}",
        f"Tipo de procedimento: {analise.get('tipo_procedimento', '?')}",
        f"Número IDEA         : {analise.get('numero_idea', '?')}",
        f"Objeto              : {analise.get('objeto', '?')}",
        f"Arquivamento viável : {'SIM' if viavel else 'NÃO'}",
        f"Hipótese            : {analise.get('hipotese_recomendada', '?')} — {analise.get('hipotese_descricao', '')}",
        f"Modelo Word         : {analise.get('modelo_word', '?')}",
        f"Justificativa       :\n  {analise.get('justificativa', analise.get('razao_nao_viavel', ''))}",
        "",
    ]
    if output_path:
        lines.append(f"Documento gerado    : {output_path.name}")
    lines += [
        "",
        "JSON COMPLETO DA ANÁLISE:",
        json.dumps(analysis, ensure_ascii=False, indent=2),
    ]

    report_path.write_text("\n".join(lines), encoding="utf-8")
    logger.info("Relatório salvo: %s", report_path)
